"""Isolated technical fixtures. These images are never production artwork."""
import copy, hashlib, importlib.util, json, sys, tempfile, unittest
from unittest.mock import patch
sys.dont_write_bytecode=True
from pathlib import Path
from PIL import Image
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2]
spec=importlib.util.spec_from_file_location('visual_builder',ROOT/'scripts/social-canva-pptx.py')
builder=importlib.util.module_from_spec(spec);spec.loader.exec_module(builder)
def fixture(folder,lesson):
    folder=Path(folder);(folder/'art').mkdir(parents=True,exist_ok=True);items=[]
    for i in range(6):
        path=folder/'art'/f'{i+1:02}.png';Image.new('RGB',(1080,1350),(240-i,238-i*2,230-i*3)).save(path)
        items.append({'file':f'art/{i+1:02}.png','sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'generation_id':f'isolated-test-{i}'})
    data={'format':'hana-social-artwork-v1','origin':'test_fixture','generation_mode':'individual_images','date':lesson['date'],'title':lesson['title'],'style':lesson.get('visuals',{}).get('style') or 'isolated fixture','preset_id':builder.PRESET['id'],'layout_reference_design_id':builder.PRESET['layout_reference_design_id'],'items':items}
    (folder/'artwork.json').write_text(json.dumps(data))
    return data
class VisualTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
        self.lesson={'date':'2026-10-27','day':9,'title':'節奏測試','visuals':{'style':'isolated fixture','preset_id':builder.PRESET['id'],'layout_reference_design_id':builder.PRESET['layout_reference_design_id']},'pages':[{'title':f'標題 {i}','body':'獨立文字內容'} for i in range(6)]}
    def tearDown(self):self.temp.cleanup()
    def test_missing_art_never_falls_back_to_shapes(self):
        self.assertIsNone(builder.artwork_files(self.lesson,self.folder))
        with self.assertRaises(ValueError):builder.create(self.lesson,self.folder/'bad.pptx')
    def test_blocked_build_keeps_daily_visual_contract(self):
        folder=self.folder/'content/instagram'/self.lesson['date'];folder.mkdir(parents=True)
        lesson={**self.lesson,'status':'editorial_ready'}
        (folder/'lesson.json').write_text(json.dumps(lesson))
        with patch.object(builder,'ROOT',self.folder),patch.dict(builder.os.environ,{'LESSON_DATE':self.lesson['date']}):builder.main()
        result=json.loads((folder/'lesson.json').read_text())
        self.assertEqual(result['visuals']['style'],self.lesson['visuals']['style'])
        self.assertEqual(result['visuals']['preset_id'],builder.PRESET['id'])
        self.assertFalse(result['visuals']['artwork_ready'])
        self.assertFalse((folder/'canva-import.pptx').exists())
    def test_test_images_cannot_be_used_in_production(self):
        fixture(self.folder,self.lesson)
        with self.assertRaises(ValueError):builder.artwork_files(self.lesson,self.folder)
    def test_six_images_and_editable_text_are_distinct(self):
        fixture(self.folder,self.lesson);art=builder.artwork_files(self.lesson,self.folder,True)
        out=self.folder/'test.pptx';builder.create(self.lesson,out,art);deck=Presentation(out)
        self.assertEqual(len(deck.slides),6);self.assertTrue(self.lesson['visuals']['fixture']);self.assertEqual(self.lesson['visuals']['human_review'],'pending')
        pictures=[]
        for i,slide in enumerate(deck.slides):
            pic=[s for s in slide.shapes if s.shape_type==13];self.assertEqual(len(pic),1);pictures.append(hashlib.sha256(pic[0].image.blob).hexdigest())
            self.assertIn(self.lesson['pages'][i]['title'],'\n'.join(s.text for s in slide.shapes if s.has_text_frame))
        self.assertEqual(len(set(pictures)),6)
        self.assertEqual(self.lesson['visuals']['layout_status'],'draft_relayout')
        self.assertEqual(self.lesson['visuals']['preset_id'],builder.PRESET['id'])
    def test_no_style_or_unrelated_preset_is_rejected(self):
        original=fixture(self.folder,self.lesson)
        self.lesson['visuals'].pop('style')
        with self.assertRaisesRegex(ValueError,'Daily style is required'):builder.artwork_files(self.lesson,self.folder,True)
        self.lesson['visuals']['style']=original['style']
        original['layout_reference_design_id']='another-template'
        (self.folder/'artwork.json').write_text(json.dumps(original))
        with self.assertRaisesRegex(ValueError,'confirmed HANA visual preset'):builder.artwork_files(self.lesson,self.folder,True)
    def test_wrong_revision_style_hash_duplicate_or_grid_is_rejected(self):
        original=fixture(self.folder,self.lesson);original_visual=copy.deepcopy(self.lesson['visuals'])
        for change in ['date','style','hash','duplicate','same_image','grid','size']:
            data=copy.deepcopy(original)
            if change=='date':data['date']='2026-10-26'
            if change=='style':self.lesson['visuals']['style']='another daily style'
            if change=='hash':data['items'][0]['sha256']='0'*64
            if change=='duplicate':data['items'][1]['generation_id']=data['items'][0]['generation_id']
            if change=='same_image':data['items'][1]['file']=data['items'][0]['file'];data['items'][1]['sha256']=data['items'][0]['sha256']
            if change=='grid':data['generation_mode']='sliced_contact_sheet'
            if change=='size':
                p=self.folder/'art/01.png';Image.new('RGB',(40,50)).save(p);data['items'][0]['sha256']=hashlib.sha256(p.read_bytes()).hexdigest()
            (self.folder/'artwork.json').write_text(json.dumps(data))
            with self.assertRaises(ValueError):builder.artwork_files(self.lesson,self.folder,True)
            self.lesson['visuals']=copy.deepcopy(original_visual)
if __name__=='__main__':
    if len(sys.argv)==3 and sys.argv[1]=='--fixture':
        folder=Path(sys.argv[2]);fixture(folder,json.loads((folder/'lesson.json').read_text()));print('Prepared isolated technical fixture; production is prohibited')
    else:unittest.main()
