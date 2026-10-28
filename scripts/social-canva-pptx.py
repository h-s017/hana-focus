#!/usr/bin/env python3
"""Create an editable PowerPoint source for Canva's officially supported REST import.
All visible text remains real PPTX text boxes. No credentials or customer records.
"""
import json
import hashlib
import os
import re
import sys
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path

from pptx import Presentation
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
PRESET=json.loads((ROOT/'social-editor/visual-preset.json').read_text(encoding='utf-8'))
def rgb(value):
    return RGBColor.from_string(value.lstrip("#"))
def label(slide,value,left,top,width,height,size,color="#f0eee8",font="Georgia",bold=False):
    shape=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height))
    tf=shape.text_frame
    tf.clear()
    tf.word_wrap=True
    tf.margin_top=Inches(.03)
    tf.margin_bottom=0
    tf.margin_left=0
    tf.margin_right=0
    for index,line in enumerate(str(value).split("\n")):
        p=tf.paragraphs[0] if index==0 else tf.add_paragraph()
        p.alignment=PP_ALIGN.LEFT
        p.space_after=Pt(6)
        run=p.add_run()
        run.text=line
        run.font.name=font
        run.font.bold=bold
        run.font.size=Pt(size)
        run.font.color.rgb=rgb(color)
    return shape
def artwork_files(lesson,folder,allow_fixture=False):
    folder=Path(folder).resolve()
    file=folder / "artwork.json"
    if not file.exists():
        return None
    artwork=json.loads(file.read_text(encoding="utf-8"))
    if artwork.get("format")!="hana-social-artwork-v1" or artwork.get("date")!=lesson["date"] or artwork.get("title")!=lesson["title"] or not artwork.get("style"):
        raise ValueError("Artwork must match this exact lesson and date")
    visual=lesson.get("visuals",{})
    if not isinstance(visual.get("style"),str) or not visual["style"].strip():
        raise ValueError("Daily style is required; do not silently select a default medium")
    if visual["style"]!=artwork["style"]:
        raise ValueError("Artwork style must match this lesson revision")
    if visual.get("preset_id")!=PRESET['id'] or artwork.get("preset_id")!=PRESET['id'] or visual.get("layout_reference_design_id")!=PRESET['layout_reference_design_id'] or artwork.get("layout_reference_design_id")!=PRESET['layout_reference_design_id']:
        raise ValueError("Artwork and lesson must use the confirmed HANA visual preset")
    if artwork.get("origin") not in ("new_original_artwork","test_fixture") or artwork.get("generation_mode")!="individual_images":
        raise ValueError("New independent artwork is required; no private-library reuse or grid slicing")
    fixture=artwork.get("origin")=="test_fixture"
    if fixture and not allow_fixture:
        raise ValueError("Test artwork must never be used by the daily production job")
    items=artwork.get("items") or []
    if len(items)!=6:
        raise ValueError("Six independently prepared artwork files are required")
    result=[];seen=set();generation_ids=set()
    for item in items:
        path=(folder/item["file"]).resolve()
        if not path.is_relative_to(folder) or not path.is_file() or path.stat().st_size>8*1024*1024:
            raise ValueError("Invalid or oversized artwork file")
        digest=hashlib.sha256(path.read_bytes()).hexdigest()
        if digest!=item.get("sha256") or digest in seen:
            raise ValueError("Artwork digest mismatch or repeated background")
        seen.add(digest)
        with Image.open(path) as image:
            width,height=image.size
            if image.format not in ("PNG","JPEG","WEBP") or width<1080 or height<1350 or abs(width/height-0.8)>0.025:
                raise ValueError("Artwork must be a full resolution individual 4:5 image")
        generation_id=item.get("generation_id")
        if not isinstance(generation_id,str) or not generation_id or generation_id in generation_ids:
            raise ValueError("Artwork must come from six separate image generations")
        generation_ids.add(generation_id)
        result.append(path)
    return {"files":result,"manifest":artwork,"fixture":fixture,"hashes":sorted(seen)}

def create(lesson,output,artwork=None):
    slides=lesson.get("pages") or []
    if len(slides)!=6 or not all(isinstance(x.get("title"),str) and isinstance(x.get("body"),str) for x in slides):
        raise ValueError("Expected exactly six real text pages")
    if not artwork:
        raise ValueError("No shape-only or dark placeholder deck may substitute for missing artwork")
    deck=Presentation();deck.slide_width=Inches(8);deck.slide_height=Inches(10)
    for i,part in enumerate(slides):
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb("#F6F3EB")
        slide.shapes.add_picture(str(artwork["files"][i]),0,0,width=deck.slide_width,height=deck.slide_height)
        label(slide,"HANA SCENT ARTIST",.67,.40,6.6,.33,11,"#62655e",font="Arial")
        label(slide,"SCENT & RHYTHM" if "節奏" in lesson["title"] else "OLFACTORY NOTES",.67,.95,6.5,.3,10,"#74786c",font="Arial")
        label(slide,part["title"],.65,1.38,6.7,1.25,34 if i==0 else 30,"#303b38",font="Noto Serif CJK TC")
        body=part["body"]
        label(slide,body,.67,2.98 if i==0 else 2.36,6.6,1.1,17,"#45514a",font="Noto Serif CJK TC")
        label(slide,"HANA 嗅覺微課 · Day "+str(lesson["day"]).zfill(3),.67,9.47,5.4,.3,10,"#5f6b60",font="Noto Sans CJK TC")
        label(slide,f"{i+1:02d} / 06",6.25,9.44,1.1,.35,12,"#45514a",font="Georgia")
    output.parent.mkdir(parents=True,exist_ok=True);deck.save(output)
    lesson["visuals"]={**lesson.get("visuals",{}),"artwork_ready":True,"image_count":6,"style":artwork["manifest"]["style"],"asset_hashes":artwork["hashes"],"pptx_sha256":hashlib.sha256(output.read_bytes()).hexdigest(),"fixture":artwork["fixture"],"layout_status":"draft_relayout","human_review":"pending","publication_authorized":False}
    return len(deck.slides)

def main():
    date=os.getenv("LESSON_DATE") or datetime.now(ZoneInfo("Asia/Taipei")).strftime("%Y-%m-%d")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}",date):
        raise SystemExit("Invalid LESSON_DATE")
    path=ROOT / "content" / "instagram" / date / "lesson.json"
    lesson=json.loads(path.read_text(encoding="utf-8"))
    if lesson.get("status")!="editorial_ready" or lesson.get("date")!=date:
        print("SKIP: no validated editorial for "+date+"; no fake Canva design created")
        return
    output=path.parent / "canva-import.pptx"
    artwork=artwork_files(lesson,path.parent,allow_fixture=os.getenv("HANA_ALLOW_TEST_ARTWORK")=="1")
    if not artwork:
        lesson["visuals"]={**lesson.get("visuals",{}),"artwork_ready":False,"image_count":0,"human_review":"not_started","reason":"missing_original_artwork"}
        path.write_text(json.dumps(lesson,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        print("BLOCKED: original artwork is missing; copy is retained, no placeholder design created")
        return
    count=create(lesson,output,artwork)
    path.write_text(json.dumps(lesson,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    check=Presentation(output)
    assert count==len(check.slides)==6
    assert check.slide_width*5==check.slide_height*4
    print("PASS: editable 6-page Canva PowerPoint source: "+str(output.relative_to(ROOT)))
    print("DRAFT: custom PowerPoint typography is not verified against the native Canva reference")
if __name__=="__main__":
    main()
